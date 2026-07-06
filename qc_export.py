#!/usr/bin/env python3
"""质控上传表填充：按模板写入数据行并打包 zip。"""
import io
import json
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
TEMPLATE_DIR = ROOT / '质控模板'


def fill_workbook(template_path: Path, rows: list) -> bytes:
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, row in enumerate(rows, start=2):
        ws.cell(i, 1, row.get('code', ''))
        ws.cell(i, 2, row.get('month', ''))
        ws.cell(i, 3, row.get('day', ''))
        ws.cell(i, 4, row.get('seq', 1))
        ws.cell(i, 5, row.get('batch', ''))
        val = row.get('value', '')
        if val is not None and val != '':
            try:
                ws.cell(i, 6, float(val))
            except (TypeError, ValueError):
                ws.cell(i, 6, val)
        ws.cell(i, 7, row.get('remark', ''))
        ws.cell(i, 8, row.get('operator', ''))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_zip(payload: dict) -> bytes:
    year = payload.get('year')
    month = payload.get('month')
    tables = payload.get('tables') or []
    config_path = ROOT / 'qc-export-config.json'
    config = {}
    if config_path.exists():
        config = json.loads(config_path.read_text(encoding='utf-8'))
    template_map = {t['id']: t.get('template', '') for t in config.get('tables', [])}

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for table in tables:
            tid = table.get('id', '')
            template_name = table.get('template') or template_map.get(tid, '')
            if not template_name:
                continue
            template_path = TEMPLATE_DIR / template_name
            if not template_path.exists():
                raise FileNotFoundError(f'模板不存在: {template_name}')
            rows = table.get('rows') or []
            xlsx_bytes = fill_workbook(template_path, rows)
            out_name = template_name.replace('_直接上传', f'_{year}-{month:02d}')
            zf.writestr(out_name, xlsx_bytes)
    return zip_buf.getvalue()